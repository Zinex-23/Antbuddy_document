#!/usr/bin/env python3
"""Build the reviewed Template functional-requirement workbook from the draft."""

from copy import copy
from math import ceil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "FR_template.xlsx"
OUTPUT = ROOT / "FR_template_completed.xlsx"


FRS = {
    "FR-TPL-001": {
        "title": "FR-TPL-001: Xem và quản lý danh sách mẫu",
        "description": "Cho phép người dùng xem, tìm kiếm, lọc và quản lý các Mẫu cấu hình AntBot trong tổ chức. Mẫu là gói cấu hình tái sử dụng; Mẫu không phải tin nhắn gửi khách và không tự tác động đến bất kỳ Page nào.",
        "actors": "Admin/Quản lý AntBot; Chủ sở hữu Mẫu; người được cấp quyền Xem, Áp dụng, Sao chép hoặc Quản lý phiên bản.",
        "pre": "Người dùng đã đăng nhập, thuộc tổ chức và có ít nhất quyền Xem Mẫu.",
        "trigger": "Người dùng mở AntBot → Mẫu cấu hình.",
        "main": "1. Hệ thống hiển thị hai phạm vi: Mẫu hệ thống và Mẫu của tổ chức/người dùng.\n2. Mỗi Mẫu hiển thị: tên, mô tả, phiên bản hiện hành, trang/kênh nguồn, chủ sở hữu, thời điểm cập nhật, trạng thái và số lần áp dụng.\n3. Người dùng tìm theo tên/mô tả; lọc theo trạng thái, kênh nguồn, trang nguồn và chủ sở hữu.\n4. Người dùng mở Xem nội dung để kiểm tra manifest, dependency, phiên bản và lịch sử áp dụng.\n5. Hệ thống hiển thị đúng hành động theo quyền: Tạo mẫu, Tạo phiên bản, Áp dụng, Sao chép, Lưu trữ, Khôi phục hoặc Xóa bản nháp chưa từng dùng.\n6. Mẫu ở trạng thái Lưu trữ chỉ được xem lịch sử/khôi phục, không được áp dụng mới.",
        "post": "Danh sách/chi tiết Mẫu được hiển thị theo đúng tổ chức và quyền. Việc xem không thay đổi Mẫu hoặc cấu hình Page.",
        "alternate": "AF-1: Chưa có Mẫu → hiển thị empty state và CTA Tạo mẫu.\nAF-2: Không có kết quả → giữ nguyên bộ lọc, hiển thị No result và nút Xóa bộ lọc.\nAF-3: Mất quyền trong lúc mở → trả Access denied, không để lộ manifest qua UI/API.\nAF-4: Dữ liệu danh sách lỗi → hiển thị trạng thái lỗi có Retry; không hiển thị dữ liệu cache của tổ chức khác.",
        "sub": "SF-1: Xem nội dung và dependency theo FR-TPL-004. SF-2: Phiên bản/lịch sử theo FR-TPL-009.",
        "ui": "UI-TPL-01 Danh sách Mẫu; UI-TPL-02 Chi tiết Mẫu; có Loading, Empty, No result, Error và phân trang.",
        "nfr": "NFR-TPL-01: tải danh sách ≤ 3 giây P95 với 1.000 Mẫu/tổ chức. NFR-TPL-02: phân quyền phía server, tenant isolation và audit. NFR-TPL-03: tìm kiếm không phân biệt hoa/thường và khoảng trắng thừa.",
        "ac": "AC-TPL-001-01: Given có Mẫu Active/Archived; When lọc Archived; Then chỉ hiện Mẫu Archived và nút Áp dụng bị khóa.\nAC-TPL-001-02: Given người dùng chỉ có quyền Xem; When mở chi tiết; Then xem được manifest nhưng không có hành động Áp dụng/Sửa.\nAC-TPL-001-03: Given không có kết quả; When xóa bộ lọc; Then danh sách mặc định được phục hồi mà không tải lại trang.",
        "br": "BR-TPL-001: Mẫu thuộc đúng một tổ chức; dữ liệu không được truy cập chéo tenant. BR-TPL-002: trạng thái Mẫu gồm DRAFT, READY, ARCHIVED. BR-TPL-003: Mẫu ARCHIVED không thể áp dụng mới. BR-TPL-004: Mẫu hệ thống chỉ đọc; người dùng phải Sao chép nếu muốn tùy biến.",
    },
    "FR-TPL-002": {
        "title": "FR-TPL-002: Tạo mẫu từ trang nguồn",
        "description": "Cho phép người dùng khởi tạo Mẫu từ cấu hình hiện có của một Page nguồn. Page nguồn chỉ được đọc; thao tác tạo Mẫu không sửa hoặc xuất bản cấu hình nguồn.",
        "actors": "Admin/Quản lý AntBot có quyền Tạo Mẫu và quyền đọc cấu hình Page nguồn.",
        "pre": "Page nguồn thuộc cùng tổ chức, đã kết nối và người dùng có quyền đọc cấu hình Page.",
        "trigger": "Tại danh sách Mẫu, người dùng bấm Tạo mẫu.",
        "main": "1. Hệ thống mở wizard và yêu cầu chọn đúng một Page nguồn.\n2. Hiển thị tên Page, kênh, trạng thái kết nối và thời điểm cấu hình được đọc gần nhất.\n3. Người dùng nhập tên Mẫu bắt buộc (1–100 ký tự) và mô tả không bắt buộc (tối đa 500 ký tự).\n4. Hệ thống chuẩn hóa tên bằng trim/collapse khoảng trắng và so sánh không phân biệt hoa thường trong tổ chức.\n5. Khi thông tin hợp lệ, hệ thống tạo/lưu Mẫu DRAFT và chuyển sang bước Chọn nội dung.\n6. Page nguồn và các Page khác không thay đổi.",
        "post": "Có Mẫu DRAFT gắn với Page/kênh nguồn, người tạo và thời điểm khởi tạo; chưa có cấu hình nào được áp dụng sang Page đích.",
        "alternate": "AF-1: Page mất kết nối hoặc người dùng mất quyền → chặn tiếp tục, nêu rõ lý do.\nAF-2: Tên rỗng/trùng/quá dài → lỗi inline, giữ dữ liệu đã nhập.\nAF-3: Page không có cấu hình thuộc phạm vi hỗ trợ → cho xem empty state nhưng không cho hoàn tất Mẫu rỗng.\nAF-4: Đổi Page nguồn sau khi đã chọn nội dung → yêu cầu xác nhận; nếu đồng ý thì xóa toàn bộ lựa chọn/dependency cũ.",
        "sub": "SF-1: Chọn nội dung theo FR-TPL-003. SF-2: Tự lưu bản nháp khi chuyển bước hoặc đóng wizard.",
        "ui": "UI-TPL-03 Wizard bước 1 — Trang nguồn & thông tin Mẫu; lỗi hiển thị cạnh đúng trường.",
        "nfr": "NFR-TPL-02: xác thực quyền tại server và audit người tạo. NFR-TPL-04: lưu DRAFT ≤ 2 giây P95; chống double-submit.",
        "ac": "AC-TPL-002-01: Given có quyền đọc Page A; When nhập tên hợp lệ; Then tạo DRAFT gắn Page A và không đổi cấu hình Page A/B.\nAC-TPL-002-02: Given đã có tên “Chăm sóc VIP”; When nhập “  chăm sóc   vip ”; Then báo trùng tên.\nAC-TPL-002-03: Given đã chọn nội dung; When đổi nguồn; Then chỉ xóa lựa chọn sau xác nhận.",
        "br": "BR-TPL-005: Một phiên bản Mẫu có đúng một Page nguồn. BR-TPL-006: tên Mẫu duy nhất trong tổ chức sau chuẩn hóa. BR-TPL-007: không hoàn tất Mẫu rỗng. BR-TPL-008: tạo Mẫu không thay đổi trạng thái/runtime của Page nguồn.",
    },
    "FR-TPL-003": {
        "title": "FR-TPL-003: Chọn thành phần và phạm vi đóng gói",
        "description": "Cho phép người dùng chọn chính xác các cấu hình cần đóng gói. Hệ thống không mặc định sao chép toàn bộ Automation của Page nguồn.",
        "actors": "Admin/Quản lý AntBot; dịch vụ đọc cấu hình Page nguồn.",
        "pre": "Mẫu DRAFT có Page nguồn hợp lệ; người dùng có quyền đọc từng cấu hình được chọn.",
        "trigger": "Người dùng vào bước Chọn nội dung hoặc chỉnh sửa manifest của Mẫu DRAFT.",
        "main": "1. Hệ thống liệt kê theo nhóm: Menu, FAQ, Tin nhắn mở đầu, Tin nhắn mặc định, Từ khóa, Luồng tin nhắn, Kịch bản chăm sóc và Quy luật.\n2. Thẻ, Trường tùy chỉnh, Opt-in, media và tài nguyên liên quan chỉ xuất hiện khi một cấu hình tham chiếu đến chúng.\n3. Ban đầu không tự chọn tất cả; người dùng chọn từng mục, chọn cả nhóm hoặc bỏ chọn.\n4. Panel manifest cập nhật ngay số lượng và tên mục đã chọn theo từng nhóm.\n5. Hệ thống ghi nhận quan hệ tham chiếu để kiểm tra tại FR-TPL-004.\n6. Người dùng lưu lựa chọn vào DRAFT; bản chụp nội dung chỉ được cố định khi hoàn tất phiên bản.",
        "post": "DRAFT lưu danh sách mục đã chọn; Page nguồn không thay đổi và chưa có Page đích nào bị ghi dữ liệu.",
        "alternate": "AF-1: Bỏ mục đang là dependency bắt buộc → đánh dấu mục phụ thuộc lỗi và chặn hoàn tất.\nAF-2: Cấu hình nguồn thay đổi/xóa trong lúc chọn → báo stale data và yêu cầu tải lại trước khi lưu.\nAF-3: Mục ngoài phạm vi hoặc chứa secret/token → vô hiệu hóa và giải thích.\nAF-4: Người dùng không còn quyền đọc một mục → bỏ mục khỏi lựa chọn và ghi audit.",
        "sub": "SF-1: Chọn tất cả/bỏ chọn tất cả chỉ tác động các mục người dùng có quyền. SF-2: Kiểm tra dependency theo FR-TPL-004.",
        "ui": "UI-TPL-04 Bộ chọn nội dung hai cột; UI-TPL-05 Manifest cập nhật thời gian thực.",
        "nfr": "NFR-TPL-04: lưu lựa chọn ≤ 2 giây P95. NFR-TPL-02: chỉ trả dữ liệu người dùng được phép đọc. NFR-TPL-05: hỗ trợ tối thiểu 500 thành phần mà không khóa UI quá 100 ms.",
        "ac": "AC-TPL-003-01: Given Page có 3 Menu, 10 Flow, 5 Tag; When chọn 1 Menu, 2 Flow, 1 Tag; Then manifest chỉ chứa đúng 4 mục đó và dependency được liệt kê riêng.\nAC-TPL-003-02: Given chọn riêng Default Message; When hoàn tất lựa chọn; Then Mẫu có thể không chứa Menu/Welcome và preview không lỗi.\nAC-TPL-003-03: Given mục chứa access token; Then token không xuất hiện trong manifest hoặc snapshot.",
        "br": "BR-TPL-009: Chỉ đóng gói mục được chọn và dependency đã xác nhận. BR-TPL-010: không đóng gói NLU/training data, khách hàng/session, phân công agent, default skill, connection credential, webhook secret hoặc báo cáo. BR-TPL-011: trạng thái Published/Active của nguồn không phải nội dung có thể tái sử dụng.",
    },
    "FR-TPL-004": {
        "title": "FR-TPL-004: Kiểm tra phụ thuộc và xem nội dung mẫu",
        "description": "Cho phép kiểm tra manifest và đồ thị phụ thuộc trước khi hoàn tất Mẫu. Xem nội dung là xem cấu hình tĩnh, không phải chạy thử chatbot.",
        "actors": "Admin/Quản lý AntBot; dịch vụ kiểm tra tham chiếu và capability.",
        "pre": "Mẫu DRAFT có ít nhất một thành phần được chọn.",
        "trigger": "Người dùng bấm Kiểm tra liên kết hoặc thay đổi lựa chọn ở FR-TPL-003.",
        "main": "1. Hiển thị manifest theo nhóm, số lượng và chi tiết từng mục.\n2. Dựng đồ thị tham chiếu cho Menu/FAQ/Keyword/Rule/Sequence/Flow và tài nguyên phụ thuộc.\n3. Phân loại dependency: bắt buộc, có thể ánh xạ ở Page đích hoặc không được phép đóng gói.\n4. Với dependency thiếu, hiển thị đường dẫn tham chiếu và cho chọn “Thêm vào Mẫu” hoặc “Bỏ mục phụ thuộc”; không tự thêm âm thầm.\n5. Chặn Hoàn tất nếu còn broken reference, mục nguồn đã bị xóa hoặc dependency không được xử lý.\n6. Khi hợp lệ, lưu manifest/dependency graph cùng content hash của bản chụp.",
        "post": "Manifest phản ánh đúng snapshot; Mẫu chỉ chuyển bước khi đồ thị dependency khép kín hoặc có kế hoạch ánh xạ hợp lệ.",
        "alternate": "AF-1: Dependency có thể dùng tài nguyên sẵn có ở đích → giữ logical key và yêu cầu ánh xạ tại FR-TPL-006.\nAF-2: Có vòng tham chiếu hợp lệ → lưu graph nhưng vẫn áp dụng loop guard runtime.\nAF-3: Nguồn đổi sau khi kiểm tra → đánh dấu snapshot stale; người dùng chọn Chụp lại hoặc giữ snapshot cũ có version rõ ràng.\nAF-4: Media đã hết hạn/không truy cập được → chặn phiên bản READY.",
        "sub": "SF-1: Xuất manifest và danh sách dependency lỗi cho kiểm toán nội bộ.",
        "ui": "UI-TPL-05 Manifest; UI-TPL-06 Dependency checker với đường dẫn From → Needs → Resolution.",
        "nfr": "NFR-TPL-06: kiểm tra ≤ 10 giây P95 với 500 node/2.000 edge. NFR-TPL-02: audit mọi quyết định thêm/bỏ dependency.",
        "ac": "AC-TPL-004-01: Given Menu gọi Flow X; When chọn Menu nhưng bỏ X; Then hiển thị đúng đường dẫn và chặn hoàn tất.\nAC-TPL-004-02: When bấm Thêm X; Then lỗi biến mất và manifest tăng đúng 1 mục.\nAC-TPL-004-03: Given nguồn đổi sau kiểm tra; Then không dùng snapshot mới âm thầm và yêu cầu người dùng quyết định.",
        "br": "BR-TPL-012: READY không được chứa broken reference. BR-TPL-013: hệ thống không âm thầm thêm dependency. BR-TPL-014: mỗi phiên bản lưu immutable snapshot, manifest, dependency graph và content hash.",
    },
    "FR-TPL-005": {
        "title": "FR-TPL-005: Hoàn tất phiên bản và phân quyền sử dụng",
        "description": "Hoàn tất một phiên bản Mẫu bất biến và cấu hình quyền xem, áp dụng, sao chép hoặc tạo phiên bản mới trong tổ chức.",
        "actors": "Chủ sở hữu Mẫu; Admin/Quản lý AntBot; người/nhóm được cấp quyền.",
        "pre": "DRAFT có manifest hợp lệ, dependency không còn lỗi chặn và người dùng có quyền Hoàn tất.",
        "trigger": "Người dùng bấm Hoàn tất tạo mẫu.",
        "main": "1. Hệ thống hiển thị lại tên, mô tả, nguồn, manifest, dependency và cảnh báo dữ liệu nhạy cảm.\n2. Người dùng chọn phạm vi: riêng tư hoặc chia sẻ cho người/nhóm trong cùng tổ chức.\n3. Cấp độc lập các quyền View, Apply, Copy và Create version; chủ sở hữu/Admin luôn có quyền quản trị theo policy tổ chức.\n4. Hệ thống chụp snapshot bất biến, tạo version number tăng dần và chuyển phiên bản sang READY.\n5. Trạng thái runtime của nguồn (Published/Active) không được kích hoạt theo Mẫu; đây chỉ là metadata tham khảo nếu cần audit.\n6. Hoàn tất Mẫu không tự áp dụng lên Page nào.",
        "post": "Có phiên bản READY bất biến và ACL rõ ràng; không thay đổi Page nguồn/đích.",
        "alternate": "AF-1: Cấp quyền cho người ngoài tổ chức → từ chối.\nAF-2: Quyền/nguồn thay đổi trước khi hoàn tất → chặn và chạy kiểm tra lại.\nAF-3: Link nội bộ được mở bởi người không có quyền → Access denied; link không thay thế ACL.\nAF-4: Trùng thao tác hoàn tất → idempotency bảo đảm chỉ tạo một version.",
        "sub": "SF-1: Thu hồi quyền không đảo ngược cấu hình đã áp dụng trước đó.",
        "ui": "UI-TPL-07 Review & Complete; UI-TPL-08 Quản lý quyền.",
        "nfr": "NFR-TPL-02: ACL bắt buộc phía server và audit. NFR-TPL-07: snapshot mã hóa khi lưu; không lộ nội dung qua URL chia sẻ.",
        "ac": "AC-TPL-005-01: Given manifest hợp lệ; When hoàn tất; Then tạo đúng một version READY và snapshot không đổi khi nguồn đổi.\nAC-TPL-005-02: Given người B chỉ có View; When gọi Apply qua UI/API; Then bị chặn.\nAC-TPL-005-03: Given nguồn có Rule Active; When hoàn tất; Then snapshot không mang quyền tự kích hoạt Rule ở đích.",
        "br": "BR-TPL-015: phiên bản READY bất biến; sửa nội dung phải tạo version mới. BR-TPL-016: link chia sẻ không thay ACL. BR-TPL-017: Template status và component runtime status là hai miền trạng thái độc lập.",
    },
    "FR-TPL-006": {
        "title": "FR-TPL-006: Chọn trang đích và xem trước tác động",
        "description": "Cho phép chọn Page đích và tạo kế hoạch áp dụng trước khi ghi dữ liệu. Màn so sánh thể hiện “Hiện tại trên Page” và “Kết quả dự kiến sau áp dụng”, không coi trạng thái runtime trong snapshot là trạng thái của Mẫu.",
        "actors": "Người có quyền Apply Mẫu và quyền ghi cấu hình trên từng Page đích.",
        "pre": "Phiên bản Mẫu READY; Page đích cùng tổ chức, đã kết nối và người dùng có quyền ghi.",
        "trigger": "Người dùng bấm Sử dụng/Áp dụng mẫu.",
        "main": "1. Người dùng chọn một hoặc nhiều Page đích; Page thuộc nhóm đồng bộ phải được nhận diện đúng scope cấu hình dùng chung.\n2. Hệ thống đọc capability của từng kênh và snapshot hiện tại của từng Page/scope.\n3. Với mỗi mục, hệ thống xác định đề xuất: KEEP (không ghi), REUSE (ánh xạ tài nguyên tương thích sẵn có), CREATE, UPDATE_AS_DRAFT hoặc UNSUPPORTED.\n4. Người dùng mở nút View để so sánh chi tiết nội dung hiện tại với nội dung/kết quả dự kiến.\n5. Cột hiện tại hiển thị trạng thái thật Published/Draft/Active/Inactive; cột Mẫu ghi “Sẽ tạo/cập nhật dưới dạng Draft/Inactive”, không hiển thị nhầm trạng thái nguồn.\n6. Thay thế Menu mặc định hoặc cấu hình đang chạy phải được chọn rõ và xác nhận; mặc định là KEEP.\n7. Hệ thống kiểm tra dependency sau từng quyết định và lập kế hoạch riêng cho từng Page; chưa ghi dữ liệu.",
        "post": "Mỗi Page/scope có application plan kèm snapshot revision, quyết định xung đột và cảnh báo; cấu hình thực tế chưa đổi.",
        "alternate": "AF-1: Kênh không hỗ trợ mục → UNSUPPORTED và chặn mục phụ thuộc nếu gây broken reference.\nAF-2: Mất quyền/mất kết nối → loại Page khỏi kế hoạch, không ảnh hưởng Page khác.\nAF-3: Page thuộc nhóm auto-sync → cảnh báo phạm vi ảnh hưởng toàn nhóm và dùng group scope.\nAF-4: Mẫu Archived hoặc version không còn Apply permission → chặn mở/xác nhận kế hoạch.",
        "sub": "SF-1: View comparison cho từng mục. SF-2: Chọn quyết định xung đột và ánh xạ dependency.",
        "ui": "UI-TPL-09 Chọn Page đích; UI-TPL-10 Application plan; UI-TPL-11 Compare current vs expected.",
        "nfr": "NFR-TPL-06: phân tích một Page ≤ 10 giây P95 với 500 mục. NFR-TPL-02: quyền ghi kiểm tra phía server. NFR-TPL-08: compare phải accessible bằng chuột và bàn phím.",
        "ac": "AC-TPL-006-01: Given Page có Menu mặc định Published; When mở plan; Then quyết định mặc định KEEP và chưa ghi dữ liệu.\nAC-TPL-006-02: When bấm View; Then hiển thị hai cột current/expected và nút không bị disabled/lỗi JS.\nAC-TPL-006-03: Given template item có status Draft từ nguồn; Then cột Mẫu ghi “Sẽ áp dụng dạng Draft”, không mô tả đó là trạng thái đang chạy.\nAC-TPL-006-04: Given Page thuộc nhóm auto-sync; Then plan nêu rõ toàn nhóm bị ảnh hưởng.",
        "br": "BR-TPL-018: không tự ghi đè cấu hình đang chạy. BR-TPL-019: mỗi Page có plan riêng. BR-TPL-020: UNSUPPORTED không được âm thầm chuyển đổi. BR-TPL-021: trạng thái hiển thị ở phía Mẫu là expected target state, không phải source runtime state.",
    },
    "FR-TPL-007": {
        "title": "FR-TPL-007: Áp dụng mẫu vào trang đích",
        "description": "Thực thi application plan đã xác nhận, tạo ID riêng tại Page đích và nối lại toàn bộ tham chiếu mà không kích hoạt cấu hình mới.",
        "actors": "Người có quyền Apply; dịch vụ sao chép cấu hình, transaction, capability và audit.",
        "pre": "Plan FR-TPL-006 hợp lệ; Mẫu READY; quyền, kết nối và revision Page vẫn còn hiệu lực.",
        "trigger": "Người dùng xác nhận kế hoạch và bấm Áp dụng.",
        "main": "1. Kiểm tra lại Mẫu, quyền, capability, kết nối và revision của từng Page/scope ngay trước khi chạy.\n2. Tạo idempotency key cho mỗi Template version + target scope + request.\n3. Tạo backup/transaction boundary; xử lý theo dependency order: resource nền → media/message/flow → menu/FAQ/keyword/sequence/rule.\n4. Mỗi mục CREATE có ID mới tại đích; REUSE dùng ID đích; toàn bộ reference được remap sang ID đích.\n5. UPDATE_AS_DRAFT chỉ cập nhật bản nháp/version mới, không ghi đè bản Published đang phục vụ khách.\n6. Menu/FAQ/message ở Draft; Keyword/Sequence/Rule và Welcome/Default Message ở Inactive sau áp dụng.\n7. Ghi Application Run và kết quả từng mục; hiển thị liên kết mở cấu hình đích.",
        "post": "Page thành công có cấu hình/link đúng ở trạng thái an toàn. Page nguồn, Template snapshot và hội thoại đang chạy không đổi.",
        "alternate": "AF-1: Revision đích đổi sau preview → dừng Page đó và yêu cầu preview lại.\nAF-2: Lỗi giữa chừng → rollback toàn bộ thay đổi của target scope; Page khác đã thành công không bị rollback.\nAF-3: Request trùng/retry → trả lại kết quả cũ, không tạo bản sao thứ hai.\nAF-4: Dependency remap thất bại → đánh dấu FAILED, rollback và ghi đường dẫn reference lỗi.\nAF-5: Page thuộc group auto-sync → ghi đúng group scope một lần, không tạo cấu hình Page riêng bị che khuất.",
        "sub": "SF-1: Application Run lưu created/reused/kept/updated/skipped/failed, ID mapping, người chạy và thời gian.",
        "ui": "UI-TPL-12 Tiến trình và kết quả áp dụng theo Page/scope; không cho double-click tạo request trùng.",
        "nfr": "NFR-TPL-09: nguyên tử trên từng target scope hoặc rollback tương đương. NFR-TPL-10: idempotent retry. NFR-TPL-02: audit đầy đủ, không log secret.",
        "ac": "AC-TPL-007-01: Given Flow X dùng Tag VIP; When áp dụng vào B; Then Flow B trỏ Tag VIP của B, không trỏ ID nguồn.\nAC-TPL-007-02: Given lỗi ở mục giữa; Then không còn dữ liệu áp dụng dở trên target scope.\nAC-TPL-007-03: Given gửi lại cùng idempotency key; Then không tăng số mục và trả cùng Run.\nAC-TPL-007-04: Given target thuộc auto-sync group; Then cấu hình được đọc thấy khi mở mọi Page thành viên.",
        "br": "BR-TPL-022: ID và reference phải thuộc target scope. BR-TPL-023: áp dụng không đồng nghĩa xuất bản/kích hoạt. BR-TPL-024: transaction tách theo target scope. BR-TPL-025: retry không nhân đôi. BR-TPL-026: backup gần nhất phải khôi phục được trước khi publish.",
    },
    "FR-TPL-008": {
        "title": "FR-TPL-008: Kiểm tra và xuất bản cấu hình sau áp dụng",
        "description": "Cho phép kiểm tra và chủ động Publish/Activate từng cấu hình sau khi áp dụng Mẫu. Chỉ bước này mới làm cấu hình tác động tới khách hàng.",
        "actors": "Người có quyền Publish/Activate trên Page đích.",
        "pre": "Application Run thành công hoặc có các mục thành công độc lập; cấu hình mới ở Draft/Inactive.",
        "trigger": "Từ kết quả áp dụng hoặc lịch sử Run, người dùng bấm Mở cấu hình/Checklist xuất bản.",
        "main": "1. Hiển thị checklist các mục vừa áp dụng, trạng thái hiện tại và dependency.\n2. Chạy validation nội dung bắt buộc, reference, channel capability và xung đột trigger.\n3. Người dùng chọn từng mục cần Publish/Activate; hệ thống hiển thị phạm vi ảnh hưởng.\n4. Publish theo dependency order; chỉ mục được chọn và hợp lệ mới hoạt động.\n5. Cập nhật Application Run với kết quả publish từng mục, người thực hiện và thời gian.\n6. Các mục không chọn hoặc lỗi vẫn Draft/Inactive và không tác động khách.",
        "post": "Chỉ cấu hình Publish/Active thành công mới phục vụ khách; phần còn lại an toàn ở Draft/Inactive.",
        "alternate": "AF-1: Connector từ chối publish → giữ Draft/Inactive, hiển thị lỗi và Retry.\nAF-2: Thiếu dependency/xung đột → chặn nhóm liên quan, không kích hoạt một phần gây broken flow.\nAF-3: Rời màn hình → tiếp tục checklist từ Application Run.\nAF-4: Chọn khôi phục trước publish → phục hồi backup của FR-TPL-007 và đóng Run là ROLLED_BACK.",
        "sub": "SF-1: Test sandbox bằng tài khoản kiểm thử; không dùng khách thật nếu chưa có quyền. SF-2: Khôi phục backup trước publish.",
        "ui": "UI-TPL-13 Post-apply checklist; UI-TPL-14 Publish result và Activity log.",
        "nfr": "NFR-TPL-02: quyền/audit. NFR-TPL-11: lỗi publish không kích hoạt cấu hình lỗi; thông báo kết quả trong ≤ 5 giây sau phản hồi connector.",
        "ac": "AC-TPL-008-01: Given Menu Draft và Rule Inactive; When chỉ Publish Menu; Then Menu hoạt động, Rule không chạy.\nAC-TPL-008-02: Given connector lỗi; Then giữ trạng thái an toàn và cho Retry.\nAC-TPL-008-03: Given chưa publish; When khách tương tác; Then vẫn nhận cấu hình Published cũ.",
        "br": "BR-TPL-027: Apply và Publish là hai bước tách biệt. BR-TPL-028: Published version cũ tiếp tục phục vụ đến khi bản mới publish thành công. BR-TPL-029: không dùng dữ liệu khách thật cho sandbox nếu thiếu quyền.",
    },
    "FR-TPL-009": {
        "title": "FR-TPL-009: Phiên bản, lịch sử áp dụng và vòng đời mẫu",
        "description": "Cho phép truy vết phiên bản, Application Run, kết quả publish; tạo phiên bản mới, lưu trữ hoặc khôi phục Mẫu.",
        "actors": "Chủ sở hữu Mẫu; Admin/Quản lý AntBot; người có quyền Create version/Audit.",
        "pre": "Mẫu tồn tại và người dùng có quyền tương ứng.",
        "trigger": "Mở Chi tiết Mẫu → Phiên bản/Lịch sử hoặc chọn Tạo phiên bản mới/Lưu trữ/Khôi phục.",
        "main": "1. Hiển thị version number, source snapshot time, content hash, người tạo, manifest và diff với phiên bản trước.\n2. Muốn thay đổi READY phải tạo DRAFT version mới; version cũ giữ nguyên.\n3. Lịch sử Run hiển thị version, target scope, plan decision, ID mapping, kết quả apply/publish/rollback, người thực hiện và thời gian.\n4. Lưu trữ Mẫu ngăn Apply mới nhưng không xóa cấu hình đã triển khai.\n5. Khôi phục Mẫu cho phép Apply lại version READY hiện hành theo quyền.\n6. Thay đổi Mẫu/Page nguồn không tự cập nhật các Page đã áp dụng.",
        "post": "Mọi lần áp dụng có thể truy vết đến version và kết quả; Page đã áp dụng không tự đổi theo version mới.",
        "alternate": "AF-1: Page nguồn bị xóa/mất quyền → vẫn xem snapshot theo ACL Mẫu, không thể chụp lại nguồn.\nAF-2: Mẫu đang được dùng → Archive chỉ chặn Run mới, không xóa dữ liệu đích.\nAF-3: Apply version Archived → chặn cả UI/API.\nAF-4: Xóa Mẫu → chỉ cho xóa DRAFT chưa có Run; Mẫu có history dùng Archive.",
        "sub": "SF-1: Export audit nếu role được phép. SF-2: Mở lại post-apply checklist từ một Run.",
        "ui": "UI-TPL-15 Version history/diff; UI-TPL-16 Application Run history; UI-TPL-02 Detail.",
        "nfr": "NFR-TPL-02: audit append-only và quyền xem. NFR-TPL-12: snapshot READY/content hash không thay đổi. NFR-TPL-13: lịch sử giữ theo retention policy tổ chức.",
        "ac": "AC-TPL-009-01: Given B đã áp dụng v1; When tạo v2; Then B vẫn giữ v1 đến khi có Run v2.\nAC-TPL-009-02: Given Mẫu Archived; When Apply qua API; Then bị chặn và không tạo Run.\nAC-TPL-009-03: Given Run thành công; Then history truy ra version, target, quyết định, ID mapping và publish result.\nAC-TPL-009-04: Given DRAFT chưa từng dùng; Then được xóa; READY/có Run chỉ được Archive.",
        "br": "BR-TPL-030: READY bất biến; version number tăng tuần tự trong Mẫu. BR-TPL-031: không tự đồng bộ Template/source → target. BR-TPL-032: history phải truy vết version và target scope. BR-TPL-033: chỉ DRAFT chưa có Run được xóa vĩnh viễn.",
    },
}


def estimated_height(text: str, width_chars: int = 112) -> float:
    lines = 0
    for paragraph in str(text).splitlines() or [""]:
        lines += max(1, ceil(len(paragraph) / width_chars))
    return max(32.0, min(230.0, lines * 17.0 + 10.0))


def write_fr_sheet(ws, fr):
    ws["A1"] = fr["title"]
    ws["A2"] = "ANTBOT · ĐẶC TẢ YÊU CẦU CHỨC NĂNG · MẪU CẤU HÌNH · BẢN 1.0"
    rows = [
        (4, "Mô tả", "description"),
        (5, "Đối tượng liên quan", "actors"),
        (6, "Pre-conditions", "pre"),
        (7, "Điều kiện kích hoạt", "trigger"),
        (8, "Luồng xử lý chính", "main"),
        (9, "Post-condition", "post"),
        (10, "Luồng thay thế", "alternate"),
        (11, "Sub-flow", "sub"),
        (12, "Giao diện hệ thống", "ui"),
        (13, "Yêu cầu phi chức năng", "nfr"),
        (14, "AC tương ứng", "ac"),
        (15, "BR tương ứng", "br"),
    ]
    for row, label, key in rows:
        ws.cell(row, 1, label)
        ws.cell(row, 2, fr[key])
        ws.cell(row, 2).alignment = copy(ws["B4"].alignment)
        ws.cell(row, 2).alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True
        )
        ws.row_dimensions[row].height = estimated_height(fr[key])
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = "A1:B15"


def build_reference_sheet(wb):
    if "Quy ước & ma trận" in wb.sheetnames:
        del wb["Quy ước & ma trận"]
    ws = wb.create_sheet("Quy ước & ma trận", 1)
    dark = PatternFill("solid", fgColor="003C90")
    pale = PatternFill("solid", fgColor="DDE7F6")
    light = PatternFill("solid", fgColor="EDF3FB")
    white_font = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
    normal = Font(name="Times New Roman", size=12, color="111111")
    bold = Font(name="Times New Roman", size=12, bold=True, color="111111")
    thin = Side(style="thin", color="8A99AD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:D1")
    ws["A1"] = "QUY ƯỚC TRẠNG THÁI VÀ MA TRẬN XỬ LÝ MẪU"
    ws["A1"].fill = dark
    ws["A1"].font = white_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    sections = [
        (3, "Vòng đời Mẫu", [
            ("DRAFT", "Đang soạn; được sửa/xóa nếu chưa có Run.", "Không được áp dụng."),
            ("READY", "Snapshot bất biến, dependency hợp lệ.", "Được preview/apply theo ACL."),
            ("ARCHIVED", "Ngừng sử dụng mới; vẫn giữ version/history.", "Không được tạo Application Run mới."),
        ]),
        (9, "Quyết định trong Application Plan", [
            ("KEEP", "Giữ nguyên cấu hình hiện tại.", "Không ghi dữ liệu."),
            ("REUSE", "Dùng tài nguyên tương thích đã có ở đích.", "Ánh xạ reference sang ID đích."),
            ("CREATE", "Tạo cấu hình mới.", "ID mới; Draft/Inactive."),
            ("UPDATE_AS_DRAFT", "Cập nhật vào bản nháp/version mới.", "Published hiện tại tiếp tục chạy."),
            ("UNSUPPORTED", "Kênh đích không hỗ trợ.", "Không áp dụng; kiểm tra dependency liên quan."),
        ]),
        (17, "Nguyên tắc trạng thái khi so sánh", [
            ("Cột hiện tại", "Hiển thị trạng thái thật trên Page.", "Published/Draft/Active/Inactive."),
            ("Cột Mẫu", "Hiển thị kết quả dự kiến sau Apply.", "Sẽ tạo/cập nhật dạng Draft/Inactive."),
            ("Snapshot nguồn", "Có thể lưu trạng thái nguồn để audit.", "Không được dùng để tự kích hoạt ở đích."),
        ]),
        (23, "Phạm vi cấu hình", [
            ("Được hỗ trợ", "Menu, FAQ, Welcome, Default Message, Keyword, Flow, Sequence, Rule và dependency được xác nhận.", "Theo capability từng kênh."),
            ("Không sao chép", "NLU/training data, khách/session, agent assignment, default skill, credential, token/secret, report.", "Luôn loại khỏi snapshot."),
        ]),
    ]
    for start, title, items in sections:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=4)
        c = ws.cell(start, 1, title)
        c.fill = dark
        c.font = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center")
        headers = ["Giá trị/Phạm vi", "Ý nghĩa", "Hành vi bắt buộc", "FR liên quan"]
        for col, value in enumerate(headers, 1):
            cell = ws.cell(start + 1, col, value)
            cell.fill = pale
            cell.font = bold
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for idx, item in enumerate(items, start + 2):
            values = [*item, "FR-TPL-005/006/007/008"]
            for col, value in enumerate(values, 1):
                cell = ws.cell(idx, col, value)
                cell.font = normal
                cell.border = border
                cell.fill = light if idx % 2 == 0 else PatternFill(fill_type=None)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[idx].height = estimated_height(" ".join(item), 75)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 28
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def main():
    wb = load_workbook(SOURCE)
    overview = wb["Danh sách FR"]
    if "A2:D2" not in [str(r) for r in overview.merged_cells.ranges]:
        overview.merge_cells("A2:D2")
    overview["A2"] = "Phiên bản 1.0 · Trạng thái: Đã rà soát · Ngày 30/09/2026"
    overview["A2"].font = Font(name="Times New Roman", size=12, italic=True, color="1D2D45")
    overview["A2"].alignment = Alignment(horizontal="center", vertical="center")
    overview.row_dimensions[2].height = 22
    overview["A16"] = "Mẫu là snapshot cấu hình có phiên bản để tái sử dụng trong cùng tổ chức. Mẫu không phải runtime configuration và không mang trạng thái Published/Active từ Page nguồn sang Page đích."
    overview["A17"] = "Phạm vi hỗ trợ: Menu, FAQ, Welcome, Default Message, Keyword, Flow, Sequence, Rule và dependency đã xác nhận. Loại trừ NLU/training data, dữ liệu khách, phân công agent, credential/token/secret và báo cáo."
    overview["A18"] = "Luồng chuẩn: chọn nguồn → chọn nội dung → kiểm tra dependency → tạo version READY → preview tác động → áp dụng Draft/Inactive → kiểm tra → Publish/Activate có chủ đích."
    for row in range(4, 13):
        target = overview.cell(row, 4).value
        overview.cell(row, 1).hyperlink = f"#'{target}'!A1"
        overview.cell(row, 1).style = "Hyperlink"
    overview.freeze_panes = "A4"
    overview.sheet_view.showGridLines = False

    for sheet_name, fr in FRS.items():
        ws = wb[sheet_name]
        write_fr_sheet(ws, fr)
        ws["A1"].hyperlink = "#'Danh sách FR'!A1"

    build_reference_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
